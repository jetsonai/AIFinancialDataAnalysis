"""예제 09 금융 업무보고서·기안문 작성 — 검증된 통계 + AI 문장 → Word 문서.

같은 분석 결과를 문서 종류만 바꿔 다시 만듭니다.
    from exercises import ex09_report
    ex09_report.DOC_TYPE = "기안문"
"""
from datetime import date

import pandas as pd
from docx.shared import Cm, Pt

from exercises.common import (DATA_NOTE, add_table, add_text, check_message, check_numbers,
                              load_customers, make_chart_files, new_document, output_path, overview,
                              rate_table, split_sections, summary_json, to_records)

DOC_TYPE = "보고서"            # "보고서" 또는 "기안문"
DIMENSIONS = ["연령대", "잔액구간", "대출상태", "과거캠페인", "직업군"]
REPORT_SECTIONS = ["주요 현황", "검토 의견", "향후 추진 방향"]
MEMO_SECTIONS = ["제목", "본문"]
CAUTIONS = [
    "이 문서는 공개된 해외 은행 캠페인 데이터(UCI Bank Marketing)를 분석한 교육용 초안이다.",
    "한국 저축은행의 실제 영업전략으로 확정하거나, 특정 고객의 가입 가능성을 단정하는 자료로 쓰지 않는다.",
    "고객군 간 차이는 함께 나타난 경향이며 원인을 증명하지 않는다. 고객수가 적은 집단은 해석에 주의한다.",
    "AI가 작성한 문장은 담당자가 원본 표와 대조해 검토한 뒤 사용한다.",
]


def highlights(tables):
    """기준마다 가입률이 가장 높은·낮은 고객군 (Python이 직접 계산)."""
    out = {}
    for dim, t in tables.items():
        hi, lo = t.loc[t["가입률(%)"].idxmax()], t.loc[t["가입률(%)"].idxmin()]
        out[dim] = {"최고": f"{hi['구분']} {hi['가입률(%)']}% ({int(hi['고객수']):,}명)",
                    "최저": f"{lo['구분']} {lo['가입률(%)']}% ({int(lo['고객수']):,}명)"}
    return out


def run_function(data_file, question):
    df = load_customers(data_file)
    tables = {dim: rate_table(df, dim) for dim in DIMENSIONS}
    return {
        "업무 요청": question,
        "문서 종류": DOC_TYPE,
        **overview(df),
        "고객군별 최고·최저": highlights(tables),
        "고객군별 가입률": {dim: to_records(t) for dim, t in tables.items()},
    }


def make_prompt(question, data_summary):
    data = {k: v for k, v in data_summary.items() if k not in ("업무 요청", "문서 종류", "고객군별 가입률")}
    if data_summary["문서 종류"] == "기안문":
        return f"""당신은 저축은행 수신기획팀 직원입니다. 아래 분석 결과를 바탕으로 내부 기안문을 작성합니다.
업무 요청: {question}

분석 결과 (Pandas 계산, 이 숫자만 사용):
{summary_json(data)}

아래 형식을 그대로 지켜 한국어로 쓰세요.
[제목]
(기안문 제목 1줄, "~ 검토(안)" 형태)
[본문]
1. 관련: (분석 근거 1줄)
2. 검토 내용: (주요 수치 2~3개를 넣어 개조식으로)
3. 건의 사항: (내부 검토를 요청하는 내용 2개, 개조식)

공문 문체(~함, ~임)로 쓰고, 데이터에 없는 숫자는 쓰지 마세요."""
    return f"""당신은 저축은행 경영기획부의 업무보고서 작성 담당자입니다.
업무 요청: {question}

분석 결과 (Pandas 계산, 이 숫자만 사용):
{summary_json(data)}

아래 세 제목을 대괄호까지 그대로 쓰고, 제목마다 내용을 작성하세요.
[주요 현황]
(전체 가입률과 고객군별 특징을 보고서 문체로 3~4문장)
[검토 의견]
(가입률이 높은 고객군·낮은 고객군의 의미와 해석상 주의점, 글머리 3개)
[향후 추진 방향]
(정기예금 가입률 향상을 위한 고객관리 방안, 글머리 3개)

보고서 문체(~로 나타남, ~할 필요가 있음)로 쓰고, 데이터에 없는 숫자는 쓰지 마세요.
연령·직업만으로 특정 고객을 고르거나 배제하는 방안은 쓰지 마세요."""


def build_report(path, title, purpose, data_summary, sections, chart_pngs, check):
    """업무보고서 Word 문서. 예제 10에서도 같은 함수를 씁니다."""
    doc = new_document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"작성일 {date.today():%Y-%m-%d} · 작성: 로컬 AI 초안 + 담당자 검토 · 교육용 문서")

    doc.add_heading("1. 분석 목적", level=1)
    doc.add_paragraph("고객 데이터의 고객군별 정기예금 가입률을 분석하여, 가입률 향상을 위한 고객관리 방안을 검토함.")
    doc.add_paragraph(f"업무 요청: {purpose}")
    doc.add_paragraph(f"전체 고객 {data_summary['전체 고객수']:,}명 중 {data_summary['전체 가입자수']:,}명이 "
                      f"정기예금에 가입하여 전체 가입률은 {data_summary['전체 가입률(%)']}%임.")
    doc.add_paragraph(DATA_NOTE)

    doc.add_heading("2. 주요 현황", level=1)
    add_text(doc, sections.get("주요 현황") or sections.get("분석 요약", ""))
    tables = data_summary["고객군별 가입률"]
    for dim in ("연령대", "대출상태"):
        if dim in tables:
            add_table(doc, pd.DataFrame(tables[dim]), caption=f"표. {dim}별 정기예금 가입률")
            doc.add_paragraph()
    for png in chart_pngs:
        doc.add_picture(png, width=Cm(15))

    doc.add_heading("3. 검토 의견", level=1)
    add_text(doc, sections.get("검토 의견") or sections.get("고객군별 특징", ""))
    doc.add_heading("4. 향후 추진 방향", level=1)
    add_text(doc, sections.get("향후 추진 방향") or sections.get("마케팅 전략 초안", ""))

    doc.add_heading("5. 유의사항", level=1)
    for line in CAUTIONS:
        doc.add_paragraph(line, style="List Bullet")
    doc.add_paragraph(check_message(check))
    doc.save(path)
    return str(path)


def build_memo(path, data_summary, sections):
    """내부 기안문 Word 문서."""
    doc = new_document()
    p = doc.add_paragraph("○○저축은행 (교육용)")
    p.runs[0].font.size, p.runs[0].bold = Pt(16), True
    doc.add_paragraph("수신  내부결재\n(경유)")
    subject = sections.get("제목", "").strip().splitlines()[0].strip() if sections.get("제목", "").strip() else ""
    if subject.startswith("(") and subject.endswith(")"):          # "(제목)"처럼 괄호로 감싼 경우만 벗기기
        subject = subject[1:-1].strip()
    doc.add_paragraph(f"제목  {subject or '정기예금 가입률 향상을 위한 고객관리 방안 검토(안)'}")
    doc.add_paragraph("")
    add_text(doc, sections.get("본문", ""))
    doc.add_paragraph("")
    doc.add_paragraph("붙임  1. 고객별_예금가입_분석.xlsx 1부.\n        2. 정기예금_가입률_향상_전략보고서.docx 1부.  끝.")
    sign = doc.add_table(rows=2, cols=3)
    sign.style = "Table Grid"
    for i, name in enumerate(["기안자", "검토자", "결재권자"]):
        sign.rows[0].cells[i].text = name
    doc.add_paragraph("")
    doc.add_paragraph(check_message(check_numbers(sections.get("본문", ""), data_summary)))
    doc.save(path)
    return str(path)


def show_result(data_summary, ai_text):
    check = check_numbers(ai_text, data_summary)
    if data_summary["문서 종류"] == "기안문":
        path = build_memo(output_path("내부기안문_초안.docx"), data_summary, split_sections(ai_text, MEMO_SECTIONS))
        return {"files": [path], "message": f"{ai_text}\n\n{check_message(check)}"}

    overall = data_summary["전체 가입률(%)"]
    tables = data_summary["고객군별 가입률"]
    pngs = []
    for dim in ("연령대", "잔액구간"):
        pngs.append(make_chart_files(pd.DataFrame(tables[dim]), dim, "가입률(%)", "막대", overall,
                                     f"ex09_{dim}_가입률")[0])
    path = build_report(output_path("정기예금_가입률_향상_전략보고서.docx"),
                        "정기예금 가입률 향상을 위한 고객관리 방안", data_summary["업무 요청"],
                        data_summary, split_sections(ai_text, REPORT_SECTIONS), pngs, check)
    return {"files": [path] + pngs, "message": f"{ai_text}\n\n{check_message(check)}"}
