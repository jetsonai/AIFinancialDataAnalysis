"""예제 06~10이 함께 쓰는 금융업무 도구 모음.

정량 계산(가입률 집계), 차트, Excel·Word 파일 만들기처럼 '정답이 정해진 일'은 모두 여기서
Python(Pandas 등)이 처리합니다. AI(LLM)는 이 결과를 해석하고 문장을 쓰는 일만 맡습니다.
"""
import json
import re
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd

OUTPUT_DIR = Path("outputs")
TEMPLATE_FILE = Path("templates/report_template.docx")

DATA_NOTE = (
    "데이터: UCI Bank Marketing (포르투갈 은행의 2008~2010년 익명화 마케팅 기록). "
    "한국 저축은행의 고객정보가 아니며, 집단 단위 교육용 분석 자료입니다."
)

# ── 1일 차와 같은 기준으로 고객군을 나눕니다 ───────────────────────
AGE_ORDER = ["20대 이하", "30대", "40대", "50대", "60대 이상"]
BALANCE_ORDER = ["마이너스", "0유로", "1~1,000", "1,001~5,000", "5,000 초과"]
LOAN_ORDER = ["대출 없음", "주택담보만", "개인대출만", "둘 다"]
CAMPAIGN_NAME = {"success": "이전 가입", "failure": "이전 미가입", "other": "기타", "unknown": "이전 연락 없음"}
JOB_NAME = {
    "admin.": "사무직", "blue-collar": "생산·현장직", "entrepreneur": "사업가",
    "housemaid": "가사도우미", "management": "관리직", "retired": "은퇴자",
    "self-employed": "자영업", "services": "서비스직", "student": "학생",
    "technician": "기술직", "unemployed": "무직", "unknown": "미상",
}

# 분석 기준 이름 → 데이터 컬럼
DIMENSIONS = {
    "연령대": "연령대",
    "직업군": "직업",
    "잔액구간": "잔액구간",
    "대출상태": "대출상태",
    "과거캠페인": "과거캠페인",
}

# 질문 속 낱말 → 분석 기준 (AI가 아니라 정해진 목록으로 해석합니다)
KEYWORDS = {
    "연령": "연령대", "나이": "연령대",
    "직업": "직업군",
    "잔액": "잔액구간", "예치": "잔액구간",
    "대출": "대출상태",
    "캠페인": "과거캠페인", "과거": "과거캠페인",
}


def load_customers(data_file):
    """고객 CSV를 읽고 분석용 항목(가입·직업·연령대·잔액구간·대출상태·과거캠페인)을 더합니다."""
    df = pd.read_csv(data_file)
    need = {"age", "job", "balance", "housing", "loan", "poutcome", "y"}
    missing = need - set(df.columns)
    if missing:
        raise ValueError(f"고객 데이터에 필요한 컬럼이 없습니다: {sorted(missing)}")

    df["가입"] = df["y"].eq("yes")
    df["직업"] = df["job"].map(JOB_NAME).fillna(df["job"])
    df["연령대"] = pd.cut(df["age"], bins=[0, 29, 39, 49, 59, 200], labels=AGE_ORDER)
    df["잔액구간"] = pd.cut(df["balance"], bins=[-np.inf, -1, 0, 1000, 5000, np.inf], labels=BALANCE_ORDER)
    house, loan = df["housing"].eq("yes"), df["loan"].eq("yes")
    df["대출상태"] = pd.Categorical(
        np.select([house & loan, house, loan], ["둘 다", "주택담보만", "개인대출만"], default="대출 없음"),
        categories=LOAN_ORDER, ordered=True)
    df["과거캠페인"] = pd.Categorical(df["poutcome"].map(CAMPAIGN_NAME),
                                    categories=list(CAMPAIGN_NAME.values()), ordered=True)
    return df


def rate_table(df, dimension):
    """분석 기준별 고객수 · 가입자수 · 가입률(%) 표."""
    col = DIMENSIONS[dimension]
    table = (
        df.groupby(col, observed=True)["가입"]
        .agg(고객수="size", 가입자수="sum")
        .reset_index()
        .rename(columns={col: "구분"})
    )
    table["가입률(%)"] = (table["가입자수"] / table["고객수"] * 100).round(1)
    table["구분"] = table["구분"].astype(str)
    if dimension == "직업군":                       # 순서가 없는 기준은 가입률 높은 순
        table = table.sort_values("가입률(%)", ascending=False).reset_index(drop=True)
    return table


def overview(df):
    """전체 고객 수 · 가입자 수 · 가입률."""
    return {
        "전체 고객수": int(len(df)),
        "전체 가입자수": int(df["가입"].sum()),
        "전체 가입률(%)": round(float(df["가입"].mean() * 100), 1),
    }


def to_records(table):
    """표를 AI에게 넘길 수 있는 목록(JSON 형태)으로 바꿉니다."""
    return json.loads(table.to_json(orient="records", force_ascii=False))


def pick_dimensions(question, default=("연령대",)):
    """질문에 나온 낱말 순서대로 분석 기준을 고릅니다. 없으면 기본값."""
    found = []
    for pos, word in sorted((question.find(w), w) for w in KEYWORDS if w in question):
        dim = KEYWORDS[word]
        if dim not in found:
            found.append(dim)
    return found or list(default)


def summary_json(data_summary):
    return json.dumps(data_summary, ensure_ascii=False, indent=1)


def output_path(name):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    return OUTPUT_DIR / name


# ── 숫자 검증: AI 문장 속 숫자가 계산 결과에 있는지 확인 ──────────────
def _collect_numbers(obj, bag):
    if isinstance(obj, dict):
        for k, v in obj.items():
            _collect_numbers(k, bag)
            _collect_numbers(v, bag)
    elif isinstance(obj, (list, tuple)):
        for v in obj:
            _collect_numbers(v, bag)
    elif isinstance(obj, bool):
        return
    elif isinstance(obj, (int, float, np.integer, np.floating)):
        bag.add(round(float(obj), 1))
    elif isinstance(obj, str):
        for m in re.findall(r"\d[\d,]*\.?\d*", obj):
            try:
                bag.add(round(float(m.replace(",", "")), 1))
            except ValueError:
                pass


def check_numbers(ai_text, data_summary):
    """AI가 쓴 숫자 중 계산 결과(data_summary)에 없는 숫자를 찾아 '확인 필요' 목록으로 돌려줍니다.

    소수점이 있거나 %·명·건·배 단위가 붙은 숫자, 100 이상인 숫자만 검사합니다.
    (목록 번호 '1.', 연령 '60대' 같은 숫자는 건너뜁니다.)
    """
    allowed = set()
    _collect_numbers(data_summary, allowed)
    pattern = re.compile(r"(?<![\w.])(\d{1,3}(?:,\d{3})+|\d+)(\.\d+)?\s*(%p|%|명|건|배|유로)?")
    checked, unmatched = [], []
    for m in pattern.finditer(ai_text):
        whole, frac, unit = m.group(1), m.group(2) or "", m.group(3) or ""
        after = ai_text[m.end():m.end() + 1]
        if after in ("대", "번", "월", "년", "단", "차", "위", "개"):
            continue
        value = float(whole.replace(",", "") + frac)
        if not (frac or unit or value >= 100):
            continue
        text = m.group(0).strip()
        checked.append(text)
        if round(value, 1) not in allowed:
            unmatched.append(text)
    return {"검사한 숫자": len(checked), "확인 필요": sorted(set(unmatched), key=unmatched.index)}


def check_message(check):
    if not check["확인 필요"]:
        return f"✅ 숫자 검증: AI 문장 속 숫자 {check['검사한 숫자']}개가 모두 계산 결과와 일치합니다."
    return (f"⚠️ 숫자 검증: {check['검사한 숫자']}개 중 계산 결과에서 찾지 못한 숫자 → "
            f"{', '.join(check['확인 필요'])}  (AI가 직접 계산했거나 지어낸 숫자일 수 있으니 원본 표와 대조하세요)")


# ── AI 답변을 [제목] 단위로 나누기 ─────────────────────────────────
def split_sections(ai_text, names):
    """'[주요 현황]'처럼 대괄호 제목으로 나뉜 AI 답변을 {제목: 내용}으로 나눕니다.
    제목을 못 찾으면 전체를 마지막 제목 아래에 넣습니다."""
    heads = []
    for name in names:
        pat = re.compile(r"^[ \t#*\d.]*[\[【]?\s*" + re.escape(name) + r"\s*[\]】]?[ \t*:：]*$", re.MULTILINE)
        m = pat.search(ai_text)
        if m:
            heads.append((m.start(), m.end(), name))
    heads.sort()
    sections = {name: "" for name in names}
    if not heads:
        sections[names[-1]] = ai_text.strip()
        return sections
    for i, (start, end, name) in enumerate(heads):
        stop = heads[i + 1][0] if i + 1 < len(heads) else len(ai_text)
        sections[name] = ai_text[end:stop].strip()
    return sections


# ── 차트 ──────────────────────────────────────────────────────────
def set_korean_font():
    """matplotlib 한글 글꼴 설정 (Windows: 맑은 고딕)."""
    import matplotlib
    from matplotlib import font_manager
    installed = {f.name for f in font_manager.fontManager.ttflist}
    for name in ["Malgun Gothic", "AppleGothic", "NanumGothic", "Noto Sans CJK KR", "Noto Sans KR"]:
        if name in installed:
            matplotlib.rcParams["font.family"] = name
            break
    matplotlib.rcParams["axes.unicode_minus"] = False


def make_chart_files(table, dimension, metric, kind, overall_rate, stem):
    """한 개의 차트를 PNG(보고서·Excel용)와 HTML(마우스 반응 Plotly)로 저장합니다.

    metric: "가입률(%)" 또는 "고객수"   kind: "막대" / "가로막대" / "원"
    """
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import plotly.express as px

    set_korean_font()
    title = f"{dimension}별 {'정기예금 가입률(%)' if metric == '가입률(%)' else '고객 수(명)'}"
    png, html = output_path(f"{stem}.png"), output_path(f"{stem}.html")

    fig, ax = plt.subplots(figsize=(7, 4), dpi=150)
    labels, values = table["구분"].tolist(), table[metric].tolist()
    if kind == "원":
        ax.pie(values, labels=labels, autopct="%1.1f%%", startangle=90, counterclock=False)
        ax.axis("equal")
    elif kind == "가로막대":
        bars = ax.barh(labels[::-1], values[::-1], color="#3776AB")
        ax.bar_label(bars, fmt="%.1f" if metric == "가입률(%)" else "%d", padding=3, fontsize=8)
        if metric == "가입률(%)":
            ax.axvline(overall_rate, color="#E5484D", linestyle=":", label=f"전체 평균 {overall_rate}%")
            ax.legend(fontsize=8)
    else:
        bars = ax.bar(labels, values, color="#3776AB")
        ax.bar_label(bars, fmt="%.1f" if metric == "가입률(%)" else "%d", padding=3, fontsize=8)
        if metric == "가입률(%)":
            ax.axhline(overall_rate, color="#E5484D", linestyle=":", label=f"전체 평균 {overall_rate}%")
            ax.legend(fontsize=8)
        ax.tick_params(axis="x", labelsize=8, rotation=30 if len(labels) > 6 else 0)
    ax.set_title(title, fontsize=11)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    fig.tight_layout()
    fig.savefig(png)
    plt.close(fig)

    if kind == "원":
        pfig = px.pie(table, names="구분", values=metric, title=title)
    else:
        pfig = px.bar(table, x="구분" if kind == "막대" else metric, y=metric if kind == "막대" else "구분",
                      orientation="v" if kind == "막대" else "h", text=metric,
                      hover_data=["고객수", "가입자수", "가입률(%)"], title=title)
        if metric == "가입률(%)":
            line = pfig.add_hline if kind == "막대" else pfig.add_vline
            line(overall_rate, line_dash="dot", annotation_text="전체 평균")
    pfig.write_html(html, include_plotlyjs=True)
    return [str(png), str(html)]


# ── Excel ─────────────────────────────────────────────────────────
def write_excel(path, basics, tables, ai_text, title="고객군별 정기예금 가입 현황"):
    """요약 시트 + 고객군별 시트(표 스타일 · 조건부 서식 · 막대 차트)로 된 Excel 보고서."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import CellIsRule, DataBarRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    base = "맑은 고딕"
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    ws["A1"] = title
    ws["A1"].font = Font(name=base, size=16, bold=True, color="2B5B84")
    ws["A2"] = f"작성일 {date.today():%Y-%m-%d} · {DATA_NOTE}"
    ws["A2"].font = Font(name=base, size=9, color="5B6778")
    row = 4
    for key, value in basics.items():
        ws.cell(row, 1, key).font = Font(name=base, bold=True)
        cell = ws.cell(row, 2, value)
        cell.font = Font(name=base, size=12, bold=True, color="2B5B84")
        cell.number_format = "0.0" if "률" in key else "#,##0"
        row += 1
    row += 1
    ws.cell(row, 1, "AI 분석 의견 (Pandas 계산 결과를 로컬 AI가 해석한 초안 · 검토 필요)").font = Font(name=base, bold=True)
    row += 1
    for line in ai_text.splitlines():
        cell = ws.cell(row, 1, line)
        cell.font = Font(name=base)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        ws.row_dimensions[row].height = max(18, 15 * (len(line) // 70 + 1))
        row += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16

    header_fill = PatternFill("solid", fgColor="2B5B84")
    overall = basics.get("전체 가입률(%)", 0)
    for i, (dim, table) in enumerate(tables.items(), start=1):
        sh = wb.create_sheet(dim)
        sh["A1"] = f"{dim}별 정기예금 가입률"
        sh["A1"].font = Font(name=base, size=13, bold=True, color="2B5B84")
        sh["A2"] = f"초록색 = 전체 평균({overall}%)보다 높은 고객군"
        sh["A2"].font = Font(name=base, size=9, color="5B6778")
        cols = list(table.columns)
        for c, name in enumerate(cols, start=1):
            cell = sh.cell(4, c, name)
            cell.font = Font(name=base, bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for r, record in enumerate(table.itertuples(index=False), start=5):
            for c, value in enumerate(record, start=1):
                cell = sh.cell(r, c, value.item() if hasattr(value, "item") else value)
                cell.font = Font(name=base)
                if cols[c - 1] == "가입률(%)":
                    cell.number_format = "0.0"
                elif cols[c - 1] in ("고객수", "가입자수"):
                    cell.number_format = "#,##0"
        last = 4 + len(table)
        ref = f"A4:{chr(64 + len(cols))}{last}"
        tab = Table(displayName=f"T{i}", ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        sh.add_table(tab)
        rate_col = chr(64 + cols.index("가입률(%)") + 1)
        rate_ref = f"{rate_col}5:{rate_col}{last}"
        sh.conditional_formatting.add(rate_ref, DataBarRule(start_type="num", start_value=0, end_type="max",
                                                            color="9CC3E6", showValue=True))
        sh.conditional_formatting.add(rate_ref, CellIsRule(operator="greaterThan", formula=[str(overall)],
                                                           fill=PatternFill("solid", fgColor="E6F6EC"),
                                                           font=Font(name=base, bold=True, color="1F9D55")))
        for c, width in zip("ABCD", (16, 10, 10, 11)):
            sh.column_dimensions[c].width = width

        chart = BarChart()
        chart.title = f"{dim}별 가입률(%)"
        chart.y_axis.title = "가입률(%)"
        chart.legend = None
        chart.height, chart.width = 7.5, 14
        data = Reference(sh, min_col=cols.index("가입률(%)") + 1, min_row=4, max_row=last)
        cats = Reference(sh, min_col=1, min_row=5, max_row=last)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        sh.add_chart(chart, "F4")
    wb.save(path)
    return str(path)


# ── Word ──────────────────────────────────────────────────────────
def new_document():
    from docx import Document
    return Document(str(TEMPLATE_FILE)) if TEMPLATE_FILE.is_file() else Document()


def add_table(doc, table, caption=None):
    from docx.shared import Pt
    if caption:
        doc.add_paragraph(caption, style="Caption" if "Caption" in [s.name for s in doc.styles] else None)
    t = doc.add_table(rows=1, cols=len(table.columns))
    t.style = "Light Grid Accent 1" if "Light Grid Accent 1" in [s.name for s in doc.styles] else t.style
    for i, name in enumerate(table.columns):
        t.rows[0].cells[i].text = str(name)
    for record in table.itertuples(index=False):
        cells = t.add_row().cells
        for i, value in enumerate(record):
            cells[i].text = f"{value:,}" if isinstance(value, (int, np.integer)) else str(value)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    return t


def add_text(doc, text):
    """AI 문장을 문단·글머리 단위로 Word에 넣습니다 (마크다운 기호 정리)."""
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        line = re.sub(r"\*\*(.+?)\*\*", r"\1", line).lstrip("#").strip()
        if re.match(r"^[-•*]\s+", line):
            doc.add_paragraph(re.sub(r"^[-•*]\s+", "", line), style="List Bullet")
        else:
            doc.add_paragraph(line)
