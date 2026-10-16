"""보충 실습 · 종합 미션용 — 보고서 구성을 직접 바꾸는 문서 에이전트 (agent.run("ex11", ...)).

예제 06~10과 같은 4단계(입력 검증 → 계산 → 프롬프트·AI → 결과)로 실행되고,
아래 설정값만 바꾸면 보고서가 달라집니다.

    from exercises import custom_report as rpt
    rpt.SECTIONS = ["핵심 요약", "고객군별 특징", "고객 안내 문구"]   # AI가 쓸 장
    rpt.LAYOUT = [...]                                                # Word 구성 (제목 · 문단 · 표 · 차트 순서)
    rpt.REPORT_TITLE, rpt.AUTHOR = "보고서 제목", "홍길동"            # 제목 · 작성자 → 파일 이름
    rpt.CAUTIONS.append("우리 지점 유의사항")                          # Python이 쓰는 유의사항
    agent.run("ex11", DATA, "요청 문장")
"""
import re
from datetime import date
from pathlib import Path

import pandas as pd
from docx.shared import Cm

from agent_core.graph import EXERCISES
from exercises import ex09_report
from exercises.common import (add_table, add_text, check_message, check_numbers, make_chart_files, new_document,
                              output_path, split_sections, summary_json)
from exercises.ex10_workflow import run_function as _run_function

EXERCISES.setdefault("ex11", "exercises.custom_report")      # agent.run("ex11", ...)으로 부를 수 있게 등록

DIMENSIONS = ["연령대", "직업군", "잔액구간", "대출상태", "과거캠페인"]
METRICS = ["가입률(%)", "고객수"]
KINDS = ["막대", "가로막대", "원"]

# ── 1. AI가 쓸 수 있는 장(섹션) 목록 — 여기서 골라 SECTIONS에 넣습니다 ─────────
SECTION_GUIDE = {
    "핵심 요약": "업무 요청에 대한 결론 3문장",
    "주요 현황": "전체 가입률과 고객군별 특징을 보고서 문체로 3~4문장",
    "고객군별 특징": "기준마다 글머리 1개, 최고·최저 가입률 숫자 포함",
    "검토 의견": "가입률이 높은·낮은 고객군의 의미와 해석상 주의점, 글머리 3개",
    "향후 추진 방향": "정기예금 가입률 향상을 위한 고객관리 방안, 글머리 3개",
    "마케팅 전략 초안": "방안 3개, 각각 '대상 · 방법 · 확인 지표'를 한 줄로",
    "고객 안내 문구": "창구 직원이 고객에게 그대로 말할 수 있는 안내 문장 2개, 존댓말",
    "리스크 및 유의사항": "해석의 한계와 고객 보호 관점의 주의점, 글머리 2개",
    "추가 확인 데이터": "결론을 내리기 전에 더 확인할 데이터 2개와 이유",
    "기대 효과": "방안을 실행했을 때 기대되는 변화 2개 (숫자로 단정하지 말 것)",
}
DEFAULT_GUIDE = "업무 요청에 맞춰 2~3문장"      # 목록에 없는 새 장 이름을 쓰면 이 안내로 작성

# ── 2. Word 보고서 구성 — 위에서부터 이 순서대로 만들어집니다 ───────────────────
# ("큰제목", "글자")            장 제목              ("작은제목", "글자")     작은 제목
# ("문단", "글자")              Python이 쓰는 문단 (아래 자리표시 사용 가능)
#        {request} 업무 요청 · {total:,} 전체 고객수 · {subscribed:,} 가입자수 · {rate} 가입률(%)
#        {author} 작성자 · {today} 작성일 · {file} 데이터 파일
# ("AI", "장 이름")             AI가 쓴 장 하나       ("AI모두",)              SECTIONS 전부 (작은제목 + 내용)
# ("표", "기준")                고객수 · 가입자수 · 가입률 표
# ("차트", "기준", "지표", "형태")   지표: 가입률(%) · 고객수 / 형태: 막대 · 가로막대 · 원(고객수만)
# ("유의사항",)  CAUTIONS 목록   ("숫자검증",)  AI 숫자 검증 결과   ("쪽나눔",)  다음 쪽으로
DEFAULT_SECTIONS = ["주요 현황", "검토 의견", "향후 추진 방향"]
DEFAULT_LAYOUT = [
    ("큰제목", "1. 분석 목적"),
    ("문단", "고객 데이터의 고객군별 정기예금 가입률을 분석하여 가입률 향상 방안을 검토함."),
    ("문단", "업무 요청: {request}"),
    ("문단", "전체 고객 {total:,}명 중 {subscribed:,}명이 정기예금에 가입하여 전체 가입률은 {rate}%임."),
    ("큰제목", "2. 분석 결과"),
    ("표", "연령대"),
    ("차트", "연령대", "가입률(%)", "막대"),
    ("표", "대출상태"),
    ("큰제목", "3. AI 분석 의견"),
    ("AI모두",),
    ("큰제목", "4. 유의사항"),
    ("유의사항",),
    ("숫자검증",),
]
DEFAULT_TITLE = "정기예금 가입률 향상을 위한 고객관리 방안"
DEFAULT_AUTHOR = "담당자"
DEFAULT_CAUTIONS = list(ex09_report.CAUTIONS)

SECTIONS = list(DEFAULT_SECTIONS)
LAYOUT = list(DEFAULT_LAYOUT)
REPORT_TITLE = DEFAULT_TITLE
AUTHOR = DEFAULT_AUTHOR
CAUTIONS = list(DEFAULT_CAUTIONS)


def reset():
    """모든 설정을 처음 상태로 되돌립니다."""
    global SECTIONS, LAYOUT, REPORT_TITLE, AUTHOR, CAUTIONS
    SECTIONS, LAYOUT = list(DEFAULT_SECTIONS), list(DEFAULT_LAYOUT)
    REPORT_TITLE, AUTHOR, CAUTIONS = DEFAULT_TITLE, DEFAULT_AUTHOR, list(DEFAULT_CAUTIONS)
    print("보고서 설정을 처음 상태로 되돌렸습니다.")


def file_name():
    return re.sub(r'[\\/:*?"<>|]', "", f"{REPORT_TITLE}_{AUTHOR}.docx").replace(" ", "_")


# ── 확인용 도우미 ────────────────────────────────────────────────────
def show_sections():
    """고를 수 있는 AI 장 목록 (✅ = 지금 고른 장)."""
    for name, guide in SECTION_GUIDE.items():
        print(f"{'✅' if name in SECTIONS else '  '} {name:<12} {guide}")
    for name in SECTIONS:
        if name not in SECTION_GUIDE:
            print(f"✅ {name:<12} (새로 만든 장) {DEFAULT_GUIDE}")


def show_layout():
    """지금 보고서 구성을 번호와 함께 보여 줍니다. LAYOUT[번호] = (...)로 한 줄씩 바꿀 수 있습니다."""
    print(f"제목: {REPORT_TITLE}   작성자: {AUTHOR}   파일: outputs/{file_name()}")
    for i, block in enumerate(LAYOUT):
        print(f"[{i:>2}] {block}")


def check_layout():
    """구성에 쓴 이름이 맞는지 미리 확인합니다. 문제가 없으면 '구성 확인 완료'."""
    problems = []
    for i, block in enumerate(LAYOUT):
        kind = block[0]
        if kind == "AI" and block[1] not in SECTIONS:
            problems.append(f"[{i}] AI 장 '{block[1]}'이 SECTIONS에 없습니다")
        elif kind == "표" and block[1] not in DIMENSIONS:
            problems.append(f"[{i}] 표 기준 '{block[1]}' → {DIMENSIONS} 중 하나")
        elif kind == "차트":
            if len(block) != 4:
                problems.append(f"[{i}] 차트는 (\"차트\", 기준, 지표, 형태) 네 칸입니다")
            elif block[1] not in DIMENSIONS or block[2] not in METRICS or block[3] not in KINDS:
                problems.append(f"[{i}] 차트 {block[1:]} → 기준 {DIMENSIONS} · 지표 {METRICS} · 형태 {KINDS}")
        elif kind not in ("큰제목", "작은제목", "문단", "AI", "AI모두", "표", "차트", "유의사항", "숫자검증", "쪽나눔"):
            problems.append(f"[{i}] 알 수 없는 종류 '{kind}'")
    print("\n".join(problems) if problems else "구성 확인 완료 ✅")
    return not problems


def show_docx(path):
    """만들어진 Word 문서의 뼈대(제목 · 문단 · 표 · 차트 순서)를 보여 줍니다."""
    from docx import Document
    doc = Document(str(path))
    n_table = n_chart = 0
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "tbl":
            n_table += 1
            print(f"   [표 {n_table}]")
        elif tag == "p":
            text = "".join(t.text or "" for t in child.iter() if t.tag.endswith("}t")).strip()
            style = child.find(".//{*}pStyle")
            style = style.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") if style is not None else ""
            if child.findall(".//{*}drawing"):
                n_chart += 1
                print(f"   [차트 {n_chart}]")
            elif text and ("Heading" in style or "Title" in style or style.startswith("1") or style.startswith("2")):
                print(("■ " if "1" in style or "Title" in style else "  ▶ ") + text)
            elif text:
                print("     " + (text[:60] + "…" if len(text) > 60 else text))
    print(f"\n표 {n_table}개 · 차트 {n_chart}개")


# ── 에이전트 4단계에서 부르는 함수 ───────────────────────────────────────
def run_function(data_file, question):
    return _run_function(data_file, question)


def make_prompt(question, data_summary):
    data = {k: v for k, v in data_summary.items() if k not in ("업무 요청", "데이터 파일", "고객군별 가입률")}
    heads = "\n".join(f"[{name}]\n({SECTION_GUIDE.get(name, DEFAULT_GUIDE)})" for name in SECTIONS)
    return f"""당신은 저축은행 수신기획팀의 데이터 분석 담당자입니다.
보고서 제목: {REPORT_TITLE}
업무 요청: {question}

Pandas로 계산한 분석 결과입니다. 이 숫자만 사용하세요.
{summary_json(data)}

아래 {len(SECTIONS)}개 제목을 대괄호까지 그대로 쓰고, 제목마다 내용을 작성하세요.
{heads}

보고서 문체(~로 나타남, ~할 필요가 있음)로 쓰고, 데이터에 없는 숫자는 쓰지 마세요.
연령·직업만으로 특정 고객을 고르거나 배제하는 방안은 쓰지 마세요."""


def show_result(data_summary, ai_text):
    check = check_numbers(ai_text, data_summary)
    sections = split_sections(ai_text, SECTIONS)
    tables = {dim: pd.DataFrame(rows) for dim, rows in data_summary["고객군별 가입률"].items()}
    overall = data_summary["전체 가입률(%)"]
    stem = Path(data_summary["데이터 파일"]).stem
    tag = "" if stem == "bank_customers" else f"_{stem}"
    values = {"request": data_summary["업무 요청"], "total": data_summary["전체 고객수"],
              "subscribed": data_summary["전체 가입자수"], "rate": overall, "author": AUTHOR,
              "today": f"{date.today():%Y-%m-%d}", "file": data_summary["데이터 파일"]}

    doc = new_document()
    doc.add_heading(REPORT_TITLE, level=0)
    doc.add_paragraph(f"작성일 {values['today']} · 작성자 {AUTHOR} · 로컬 AI 초안 + 담당자 검토 · 교육용 문서")
    charts = []
    for block in LAYOUT:
        kind = block[0]
        if kind == "큰제목":
            doc.add_heading(block[1], level=1)
        elif kind == "작은제목":
            doc.add_heading(block[1], level=2)
        elif kind == "문단":
            try:
                doc.add_paragraph(block[1].format(**values))
            except (KeyError, ValueError, IndexError):
                doc.add_paragraph(block[1])          # 자리표시를 잘못 쓰면 글자 그대로
        elif kind == "AI":
            add_text(doc, sections.get(block[1]) or "(AI 답에서 이 장을 찾지 못함 — 다시 실행해 보세요)")
        elif kind == "AI모두":
            for name in SECTIONS:
                doc.add_heading(name, level=2)
                add_text(doc, sections.get(name) or "(AI 답에서 이 장을 찾지 못함 — 다시 실행해 보세요)")
        elif kind == "표":
            add_table(doc, tables[block[1]], caption=f"표. {block[1]}별 정기예금 가입률")
            doc.add_paragraph()
        elif kind == "차트":
            dim, metric, chart_kind = block[1], block[2], block[3]
            if chart_kind == "원" and metric != "고객수":
                metric = "고객수"                   # 원그래프는 비중(고객수)만 그립니다
            files = make_chart_files(tables[dim], dim, metric, chart_kind, overall,
                                     f"ex11_{len(charts) + 1}_{dim}{tag}")
            charts += files
            doc.add_picture(files[0], width=Cm(14))
        elif kind == "유의사항":
            for line in CAUTIONS:
                doc.add_paragraph(line, style="List Bullet")
        elif kind == "숫자검증":
            doc.add_paragraph(check_message(check))
        elif kind == "쪽나눔":
            doc.add_page_break()
    path = output_path(file_name())
    doc.save(path)
    return {"files": [str(path)] + charts, "message": f"{ai_text}\n\n{check_message(check)}"}
